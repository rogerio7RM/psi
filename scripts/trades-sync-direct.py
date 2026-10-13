#!/usr/bin/env python3
"""Download authoritative OneDrive workbook and build private R2 snapshots."""
import datetime as dt
import io
import json
import os
import pathlib
import sys
import urllib.parse
import urllib.request
from openpyxl import load_workbook

DRIVE = "51DC8D218DF0FB5C"
ITEM = "51DC8D218DF0FB5C!sa01fad4ef8b64abb9005b5a015ada62f"
SOURCE = "PrimeSphere/Operacoes Logbook.xlsx"
FIELDS = ("sourceRow","date","asset","strike","quantity","strategy","rawAmount","status")
OUT = pathlib.Path("/tmp/trades-direct")
OUT.mkdir(exist_ok=True)

def request(url, data=None, token=None):
    headers = {"Accept":"application/json"}
    if token: headers["Authorization"] = "Bearer " + token
    if data is not None: headers["Content-Type"] = "application/x-www-form-urlencoded"
    with urllib.request.urlopen(urllib.request.Request(url, data=data, headers=headers), timeout=45) as r:
        return r.read()

def get_token():
    cid = os.environ["MS_GRAPH_CLIENT_ID"]
    refresh = pathlib.Path("/tmp/trades-graph-refresh-token").read_text(encoding="utf-8").strip() if pathlib.Path("/tmp/trades-graph-refresh-token").exists() else os.environ["MS_GRAPH_REFRESH_TOKEN"]
    body = urllib.parse.urlencode({
        "client_id":cid, "grant_type":"refresh_token", "refresh_token":refresh,
        "scope":"offline_access Files.Read"
    }).encode()
    obj = json.loads(request("https://login.microsoftonline.com/consumers/oauth2/v2.0/token",body))
    return obj["access_token"]

def as_text(v):
    if v is None: return ""
    return str(v).strip()

def json_value(v):
    if isinstance(v, (dt.datetime,dt.date)): return v.date().isoformat() if isinstance(v,dt.datetime) else v.isoformat()
    if v is None: return ""
    return str(v)

token = get_token()
base = "https://graph.microsoft.com/v1.0/drives/" + DRIVE + "/items/" + ITEM
meta = json.loads(request(base,token=token))
if meta.get("name") != "Operacoes Logbook.xlsx": raise ValueError("Unexpected source file")
# Microsoft Graph returns a short-lived, preauthenticated download URL.
# Do not forward the Graph bearer token to the storage download host.
download_url = meta.get("@microsoft.graph.downloadUrl")
if not download_url or not download_url.startswith("https://"):
    raise ValueError("Microsoft Graph did not return a secure download URL")
raw = request(download_url)
wb = load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
needed = {"SECRET FLAG","Data","ATIVO","STRIKE","QTDD","OPERACAO","VALOR","OBS","RM","EB","DC"}
sheet = None
headers = None
for candidate in wb.worksheets:
    first = next(candidate.iter_rows(min_row=1,max_row=1,values_only=True),None)
    if first:
        found = {as_text(v):i for i,v in enumerate(first)}
        if needed.issubset(found):
            sheet,headers = candidate,found
            break
if sheet is None: raise ValueError("Missing expected Excel headers")
accounts = {"rm":[],"eb":[],"dc":[]}
secret = 0
for rownum, cells in enumerate(sheet.iter_rows(min_row=2,values_only=True),start=2):
    def cell(name): return cells[headers[name]] if headers[name] < len(cells) else None
    if not any(v is not None and as_text(v) for v in cells): continue
    if as_text(cell("SECRET FLAG")).upper()=="Y":
        secret+=1
        continue
    targets = [a for a in accounts if as_text(cell(a.upper())).upper()=="Y"]
    if not targets: continue
    amount = cell("VALOR")
    if isinstance(amount,bool) or not isinstance(amount,(int,float)):
        raise ValueError("Invalid VALOR on row " + str(rownum))
    date = cell("Data")
    if isinstance(date,(dt.datetime,dt.date)):
        date = date.date().isoformat() if isinstance(date,dt.datetime) else date.isoformat()
    else: date = as_text(date)
    trade = {
        "sourceRow":rownum,"date":date,"asset":json_value(cell("ATIVO")),
        "strike":json_value(cell("STRIKE")),"quantity":cell("QTDD"),
        "strategy":json_value(cell("OPERACAO")),"rawAmount":float(amount),
        "status":"closed" if as_text(cell("OBS")).upper()=="CLOSED" else "open"
    }
    if set(trade)!=set(FIELDS): raise ValueError("Unexpected snapshot keys")
    for a in targets: accounts[a].append(trade)
stamp = dt.datetime.now(dt.timezone.utc).isoformat().replace("+00:00","Z")
for a,rows in accounts.items():
    payload = {"source":SOURCE,"sourceModifiedAt":meta["lastModifiedDateTime"],
               "syncedAt":stamp,"account":a,"rows":rows}
    (OUT/(a+".json")).write_text(json.dumps(payload,ensure_ascii=False,separators=(",",":"),default=str),encoding="utf-8")
print(json.dumps({"sourceModifiedAt":meta["lastModifiedDateTime"],"counts":{a:len(r) for a,r in accounts.items()},"secretExcluded":secret}))
